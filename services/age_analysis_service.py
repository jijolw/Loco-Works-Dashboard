# =====================================================
# services/age_analysis_service.py
# LW/PER Carriage Workshop Intelligence System
# ICF 13-Year Age Condition & Condemnation Survey Service
# =====================================================

import os
import sys
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from services.erp_service import get_session, fetch_coach_meta
from services.type_wise_holding_service import map_coach_to_category, decode_family
import config

FLOOR_OVERRIDES = {
    "101102": "HCB/DESP",
    "136646": "HCB/DESP"
}

def get_area(pit):
    pit = str(pit or "").strip()
    if pit.startswith(("AS/", "HCB/", "LCB/", "LBR/", "PS/", "DM/")):
        return "Working Area"
    return "Yard"

def get_live_icf_coaches():
    sess = get_session()
    r = sess.get(f"{config.COACH_ERP_API_BASE}/locos/masters/coach-receipts", timeout=12)
    receipts = r.json() if r.status_code == 200 else []

    running = [
        c for c in receipts 
        if str(c.get("status")).strip() == "Running" 
        and not c.get("dispatchDate") 
        and not c.get("actualDispatchDate")
    ]

    coaches = []
    current_year = datetime.now().year

    for c in running:
        cno = str(c.get("coachNo")).strip()
        meta = fetch_coach_meta(cno)
        desc = meta.get("coachTypeDescription") or c.get("coachDesc") or ""
        cat = map_coach_to_category(desc, repair_type=str(c.get("repairType") or ""))
        
        # Check if ICF coaching stock
        if cat in ("CN", "GS", "SLR", "GSLRD", "CZ", "CZJ") or desc in ("CN", "GS", "SLR", "GSLRD", "CZ", "CZJ"):
            pit = FLOOR_OVERRIDES.get(cno, str(c.get("pitNum") or "").strip())
            yb_raw = meta.get("yearBuilt")
            try:
                yb = int(yb_raw) if yb_raw else None
            except:
                yb = None
            
            age = (current_year - yb) if yb else None
            is_ge2013 = (yb is not None and yb >= 2013)
            
            group_label = "CZ / CZJ" if cat in ("CZ", "CZJ", "CZRJ") else ("SLR / GSLRD" if cat in ("SLR", "GSLRD") else cat)
            rly = meta.get("divisionName") or "SR"
            div = meta.get("divisionId") or "PGT"
            
            m_hrs = c.get("finalHrs") or c.get("preSurveyHrs")
            try:
                man_hours = int(m_hrs) if m_hrs is not None else None
            except:
                man_hours = None
            
            coaches.append({
                "cno": cno,
                "desc": desc,
                "type": cat,
                "group": group_label,
                "yb": yb,
                "age": age,
                "is_ge2013": is_ge2013,
                "age_cat": ">= 2013 (<= 13 Yrs)" if is_ge2013 else "< 2013 (> 13 Yrs)",
                "survey_flag": "Retained" if is_ge2013 else "Condemnation Survey",
                "rly": rly,
                "div": div,
                "pit": pit,
                "area": get_area(pit),
                "man_hours": man_hours,
                "recd": c.get("receivedDate") or "",
                "stage": c.get("stageId")
            })

    return coaches

def get_age_analysis_summary():
    coaches = get_live_icf_coaches()
    tot_ge = sum(1 for c in coaches if c["is_ge2013"])
    tot_lt = sum(1 for c in coaches if not c["is_ge2013"])
    tot_all = len(coaches)
    tot_hrs = sum(c["man_hours"] for c in coaches if c["man_hours"] is not None)

    type_meta = [
        ("CN", "Sleeper Non-AC (WGSCN)"),
        ("GS", "Second Class General (GS)"),
        ("CZ / CZJ", "Chair Car Non-AC (CZ / CZJ)"),
        ("SLR / GSLRD", "Luggage & Brake Van (SLR/GSLRD)")
    ]

    breakdown = []
    for grp_label, grp_desc in type_meta:
        m_ge = sum(1 for c in coaches if c["group"] == grp_label and c["is_ge2013"])
        m_lt = sum(1 for c in coaches if c["group"] == grp_label and not c["is_ge2013"])
        m_tot = m_ge + m_lt
        m_pct = (m_ge / m_tot * 100) if m_tot > 0 else 0
        grp_hrs = sum(c["man_hours"] for c in coaches if c["group"] == grp_label and c["man_hours"] is not None)
        breakdown.append({
            "type": grp_label,
            "description": grp_desc,
            "ge_2013": m_ge,
            "lt_2013": m_lt,
            "total": m_tot,
            "man_hours": grp_hrs,
            "pct_ge_2013": round(m_pct, 1)
        })

    return {
        "timestamp": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "total_coaches": tot_all,
        "ge_2013_count": tot_ge,
        "ge_2013_pct": round((tot_ge / tot_all * 100) if tot_all else 0, 1),
        "lt_2013_count": tot_lt,
        "lt_2013_pct": round((tot_lt / tot_all * 100) if tot_all else 0, 1),
        "total_man_hours": tot_hrs,
        "breakdown": breakdown,
        "coaches": coaches
    }

def generate_age_analysis_excel_bytes():
    coaches = get_live_icf_coaches()
    wb = openpyxl.Workbook()

    navy_dark = "1B365D"
    blue_header = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

    fill_ge2013 = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    font_ge2013 = Font(name="Calibri", size=11, bold=True, color="1F497D")

    fill_lt2013 = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    font_lt2013 = Font(name="Calibri", size=11, bold=True, color="C00000")

    total_fill = PatternFill(start_color="B4C6E7", end_color="B4C6E7", fill_type="solid")
    total_font = Font(name="Calibri", size=11, bold=True, color="000000")

    sub_fill = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
    data_font = Font(name="Calibri", size=11)
    data_bold = Font(name="Calibri", size=11, bold=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    double_bottom_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='double', color='000000')
    )

    def setup_ws(ws):
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.4
        ws.page_margins.bottom = 0.4

    tot_ge = sum(1 for c in coaches if c["is_ge2013"])
    tot_lt = sum(1 for c in coaches if not c["is_ge2013"])
    tot_all = len(coaches)
    pct_ge = (tot_ge / tot_all * 100) if tot_all else 0
    pct_lt = (tot_lt / tot_all * 100) if tot_all else 0
    tot_man_hours = sum(c["man_hours"] for c in coaches if c["man_hours"] is not None)

    # TAB 1: Summary
    ws1 = wb.active
    ws1.title = "Type-Wise Age Summary"
    setup_ws(ws1)

    ws1.merge_cells("A1:H1")
    ws1["A1"] = f"CARRIAGE WORKSHOP - LIVE {tot_all} ICF COACHES AGE PROFILE & MAN-HOURS SUMMARY"
    ws1["A1"].font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    ws1["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 28

    ws1.merge_cells("A2:H2")
    ws1["A2"] = f"Live Position as of {datetime.now().strftime('%d/%m/%Y %H:%M')} | Age Classification: >= 2013 (<= 13 Yrs) vs Below 2013 (> 13 Yrs Condemnation Survey)"
    ws1["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")
    ws1["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[2].height = 18

    ws1.merge_cells("B4:C4")
    ws1["B4"] = "BUILT >= 2013 (<= 13 YRS OLD)"
    ws1["B4"].font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    ws1["B4"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["B4"].fill = fill_ge2013

    ws1.merge_cells("B5:C5")
    ws1["B5"] = f"{tot_ge} COACHES ({pct_ge:.1f}%)"
    ws1["B5"].font = Font(name="Calibri", size=13, bold=True, color="1F497D")
    ws1["B5"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["B5"].fill = fill_ge2013

    ws1.merge_cells("D4:E4")
    ws1["D4"] = "BUILT < 2013 (> 13 YRS - SURVEY)"
    ws1["D4"].font = Font(name="Calibri", size=11, bold=True, color="C00000")
    ws1["D4"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["D4"].fill = fill_lt2013

    ws1.merge_cells("D5:E5")
    ws1["D5"] = f"{tot_lt} COACHES ({pct_lt:.1f}%)"
    ws1["D5"].font = Font(name="Calibri", size=13, bold=True, color="C00000")
    ws1["D5"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["D5"].fill = fill_lt2013

    ws1.merge_cells("F4:G4")
    ws1["F4"] = "TOTAL LOGGED MAN-HOURS"
    ws1["F4"].font = Font(name="Calibri", size=11, bold=True, color="000000")
    ws1["F4"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["F4"].fill = total_fill

    ws1.merge_cells("F5:G5")
    ws1["F5"] = f"{tot_man_hours:,} HRS"
    ws1["F5"].font = Font(name="Calibri", size=13, bold=True, color="000000")
    ws1["F5"].alignment = Alignment(horizontal="center", vertical="center")
    ws1["F5"].fill = total_fill

    ws1.row_dimensions[4].height = 20
    ws1.row_dimensions[5].height = 24

    headers1 = [
        "S.No", "Coach Type", "Type Description", 
        "Built >= 2013 (<= 13 Yrs)", "Built < 2013 (> 13 Yrs)", 
        "Total Coaches", "Total Man-Hours", "% Share (>= 2013)"
    ]
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=7, column=col_idx, value=h)
        cell.fill = blue_header
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border
    ws1.row_dimensions[7].height = 28

    type_meta = [
        ("CN", "Sleeper Non-AC (WGSCN)"),
        ("GS", "Second Class General (GS)"),
        ("CZ / CZJ", "Chair Car Non-AC (CZ / CZJ)"),
        ("SLR / GSLRD", "Luggage & Brake Van (SLR/GSLRD)")
    ]

    row_curr = 8
    sno = 1
    for grp_label, grp_desc in type_meta:
        m_ge = sum(1 for c in coaches if c["group"] == grp_label and c["is_ge2013"])
        m_lt = sum(1 for c in coaches if c["group"] == grp_label and not c["is_ge2013"])
        m_tot = m_ge + m_lt
        m_pct = (m_ge / m_tot * 100) if m_tot > 0 else 0
        grp_hrs = sum(c["man_hours"] for c in coaches if c["group"] == grp_label and c["man_hours"] is not None)

        ws1.cell(row=row_curr, column=1, value=sno).alignment = Alignment(horizontal="center")
        ws1.cell(row=row_curr, column=2, value=grp_label).font = data_bold
        ws1.cell(row=row_curr, column=3, value=grp_desc)
        
        c_ge = ws1.cell(row=row_curr, column=4, value=m_ge)
        c_ge.alignment = Alignment(horizontal="center")
        c_ge.font = font_ge2013
        c_ge.fill = fill_ge2013

        c_lt = ws1.cell(row=row_curr, column=5, value=m_lt)
        c_lt.alignment = Alignment(horizontal="center")
        c_lt.font = font_lt2013
        c_lt.fill = fill_lt2013

        c_tot = ws1.cell(row=row_curr, column=6, value=m_tot)
        c_tot.alignment = Alignment(horizontal="center")
        c_tot.font = data_bold

        c_hrs = ws1.cell(row=row_curr, column=7, value=grp_hrs if grp_hrs > 0 else "-")
        c_hrs.alignment = Alignment(horizontal="center")
        c_hrs.font = data_bold
        if isinstance(c_hrs.value, int):
            c_hrs.number_format = "#,##0"

        c_pct = ws1.cell(row=row_curr, column=8, value=f"{m_pct:.1f}%")
        c_pct.alignment = Alignment(horizontal="center")
        c_pct.font = data_bold

        for c_idx in range(1, 9):
            ws1.cell(row=row_curr, column=c_idx).border = thin_border
        ws1.row_dimensions[row_curr].height = 22
        row_curr += 1
        sno += 1

    ws1.cell(row=row_curr, column=2, value="TOTAL").font = total_font
    ws1.cell(row=row_curr, column=3, value="All Live ICF Coaches").font = total_font
    ws1.cell(row=row_curr, column=4, value=tot_ge).font = total_font
    ws1.cell(row=row_curr, column=4).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_curr, column=5, value=tot_lt).font = total_font
    ws1.cell(row=row_curr, column=5).alignment = Alignment(horizontal="center")
    ws1.cell(row=row_curr, column=6, value=tot_all).font = total_font
    ws1.cell(row=row_curr, column=6).alignment = Alignment(horizontal="center")
    
    tot_hrs_cell = ws1.cell(row=row_curr, column=7, value=tot_man_hours)
    tot_hrs_cell.font = total_font
    tot_hrs_cell.alignment = Alignment(horizontal="center")
    tot_hrs_cell.number_format = "#,##0"

    ws1.cell(row=row_curr, column=8, value=f"{pct_ge:.1f}%").font = total_font
    ws1.cell(row=row_curr, column=8).alignment = Alignment(horizontal="center")

    for c_idx in range(1, 9):
        cell = ws1.cell(row=row_curr, column=c_idx)
        cell.fill = total_fill
        cell.border = double_bottom_border
    ws1.row_dimensions[row_curr].height = 25

    ws1.column_dimensions["A"].width = 8
    ws1.column_dimensions["B"].width = 16
    ws1.column_dimensions["C"].width = 32
    ws1.column_dimensions["D"].width = 24
    ws1.column_dimensions["E"].width = 24
    ws1.column_dimensions["F"].width = 16
    ws1.column_dimensions["G"].width = 18
    ws1.column_dimensions["H"].width = 18

    # TAB 2: >= 2013
    ws2 = wb.create_sheet(f"2013 & Above ({tot_ge})")
    setup_ws(ws2)
    ws2.merge_cells("A1:I1")
    ws2["A1"] = f"LIVE ICF COACHES - MANUFACTURED IN 2013 & ABOVE (<= 13 YRS OLD) [TOTAL: {tot_ge} COACHES]"
    ws2["A1"].font = Font(name="Calibri", size=13, bold=True, color="1F497D")
    ws2["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws2.row_dimensions[1].height = 26

    headers2 = ["Sl.No", "Coach Type", "Coach No", "Owning Rly", "Division", "Year Built", "Age (Yrs)", "Man-Hours", "Current Location / Pit"]
    for col_idx, h in enumerate(headers2, 1):
        c = ws2.cell(row=3, column=col_idx, value=h)
        c.fill = blue_header
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws2.row_dimensions[3].height = 25

    ge_coaches = [c for c in coaches if c["is_ge2013"]]
    ge_coaches.sort(key=lambda x: (x["group"], -(x["yb"] or 0), x["cno"]))

    row_t2 = 4
    sno_t2 = 1
    current_grp = None
    for c in ge_coaches:
        if c["group"] != current_grp:
            current_grp = c["group"]
            grp_cnt = sum(1 for x in ge_coaches if x["group"] == current_grp)
            ws2.merge_cells(start_row=row_t2, start_column=1, end_row=row_t2, end_column=9)
            sh = ws2.cell(row=row_t2, column=1, value=f"{current_grp} - {grp_cnt} COACHES (MANUFACTURED >= 2013)")
            sh.fill = sub_fill
            sh.font = Font(name="Calibri", size=11, bold=True, color="1B365D")
            sh.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for c_idx in range(1, 10):
                ws2.cell(row=row_t2, column=c_idx).border = thin_border
            ws2.row_dimensions[row_t2].height = 22
            row_t2 += 1

        hrs_val = c["man_hours"] if c["man_hours"] is not None else "-"
        vals = [sno_t2, c["type"], c["cno"], c["rly"], c["div"], c["yb"], c["age"], hrs_val, f"{c['pit']} ({c['area']})"]
        for col_idx, v in enumerate(vals, 1):
            cell = ws2.cell(row=row_t2, column=col_idx, value=v)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 2, 4, 5, 6, 7, 8):
                cell.alignment = Alignment(horizontal="center")
            if col_idx == 3:
                cell.alignment = Alignment(horizontal="center")
                cell.font = data_bold
            if col_idx == 6:
                cell.fill = fill_ge2013
                cell.font = font_ge2013
            if col_idx == 8 and isinstance(v, int):
                cell.number_format = "#,##0"
                cell.font = data_bold
        ws2.row_dimensions[row_t2].height = 20
        row_t2 += 1
        sno_t2 += 1

    ws2.cell(row=row_t2, column=2, value="TOTAL").font = total_font
    ws2.cell(row=row_t2, column=3, value=f"{tot_ge} Coaches").font = total_font
    ws2.cell(row=row_t2, column=3).alignment = Alignment(horizontal="center")
    ge_hrs_tot = sum(c["man_hours"] for c in ge_coaches if c["man_hours"] is not None)
    c_hrs2 = ws2.cell(row=row_t2, column=8, value=ge_hrs_tot)
    c_hrs2.font = total_font
    c_hrs2.alignment = Alignment(horizontal="center")
    c_hrs2.number_format = "#,##0"
    for col_idx in range(1, 10):
        c = ws2.cell(row=row_t2, column=col_idx)
        c.fill = total_fill
        c.border = double_bottom_border
    ws2.row_dimensions[row_t2].height = 24

    ws2.column_dimensions["A"].width = 8
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 16
    ws2.column_dimensions["D"].width = 14
    ws2.column_dimensions["E"].width = 14
    ws2.column_dimensions["F"].width = 14
    ws2.column_dimensions["G"].width = 12
    ws2.column_dimensions["H"].width = 16
    ws2.column_dimensions["I"].width = 28

    # TAB 3: < 2013
    ws3 = wb.create_sheet(f"Below 2013 ({tot_lt})")
    setup_ws(ws3)
    ws3.merge_cells("A1:I1")
    ws3["A1"] = f"LIVE ICF COACHES - MANUFACTURED BEFORE 2013 (> 13 YRS OLD) [CONDEMNATION SURVEY] [TOTAL: {tot_lt} COACHES]"
    ws3["A1"].font = Font(name="Calibri", size=13, bold=True, color="C00000")
    ws3["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws3.row_dimensions[1].height = 26

    for col_idx, h in enumerate(headers2, 1):
        c = ws3.cell(row=3, column=col_idx, value=h)
        c.fill = blue_header
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws3.row_dimensions[3].height = 25

    lt_coaches = [c for c in coaches if not c["is_ge2013"]]
    lt_coaches.sort(key=lambda x: (x["group"], (x["yb"] or 0), x["cno"]))

    row_t3 = 4
    sno_t3 = 1
    current_grp3 = None
    for c in lt_coaches:
        if c["group"] != current_grp3:
            current_grp3 = c["group"]
            grp_cnt = sum(1 for x in lt_coaches if x["group"] == current_grp3)
            ws3.merge_cells(start_row=row_t3, start_column=1, end_row=row_t3, end_column=9)
            sh = ws3.cell(row=row_t3, column=1, value=f"{current_grp3} - {grp_cnt} COACHES (MANUFACTURED < 2013 / > 13 YRS)")
            sh.fill = sub_fill
            sh.font = Font(name="Calibri", size=11, bold=True, color="C00000")
            sh.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for c_idx in range(1, 10):
                ws3.cell(row=row_t3, column=c_idx).border = thin_border
            ws3.row_dimensions[row_t3].height = 22
            row_t3 += 1

        hrs_val = c["man_hours"] if c["man_hours"] is not None else "-"
        vals = [sno_t3, c["type"], c["cno"], c["rly"], c["div"], c["yb"], c["age"], hrs_val, f"{c['pit']} ({c['area']})"]
        for col_idx, v in enumerate(vals, 1):
            cell = ws3.cell(row=row_t3, column=col_idx, value=v)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 2, 4, 5, 6, 7, 8):
                cell.alignment = Alignment(horizontal="center")
            if col_idx == 3:
                cell.alignment = Alignment(horizontal="center")
                cell.font = data_bold
            if col_idx == 6:
                cell.fill = fill_lt2013
                cell.font = font_lt2013
            if col_idx == 8 and isinstance(v, int):
                cell.number_format = "#,##0"
                cell.font = data_bold
        ws3.row_dimensions[row_t3].height = 20
        row_t3 += 1
        sno_t3 += 1

    ws3.cell(row=row_t3, column=2, value="TOTAL").font = total_font
    ws3.cell(row=row_t3, column=3, value=f"{tot_lt} Coaches").font = total_font
    ws3.cell(row=row_t3, column=3).alignment = Alignment(horizontal="center")
    lt_hrs_tot = sum(c["man_hours"] for c in lt_coaches if c["man_hours"] is not None)
    c_hrs3 = ws3.cell(row=row_t3, column=8, value=lt_hrs_tot)
    c_hrs3.font = total_font
    c_hrs3.alignment = Alignment(horizontal="center")
    c_hrs3.number_format = "#,##0"
    for col_idx in range(1, 10):
        c = ws3.cell(row=row_t3, column=col_idx)
        c.fill = total_fill
        c.border = double_bottom_border
    ws3.row_dimensions[row_t3].height = 24

    ws3.column_dimensions["A"].width = 8
    ws3.column_dimensions["B"].width = 14
    ws3.column_dimensions["C"].width = 16
    ws3.column_dimensions["D"].width = 14
    ws3.column_dimensions["E"].width = 14
    ws3.column_dimensions["F"].width = 14
    ws3.column_dimensions["G"].width = 12
    ws3.column_dimensions["H"].width = 16
    ws3.column_dimensions["I"].width = 28

    # TAB 4: Master List
    ws4 = wb.create_sheet(f"Live {tot_all} Coaches List")
    setup_ws(ws4)
    ws4.merge_cells("A1:J1")
    ws4["A1"] = f"CARRIAGE WORKSHOP - COMPLETE LIVE {tot_all} ICF COACHES MASTER LIST WITH AGE & MAN-HOURS"
    ws4["A1"].font = Font(name="Calibri", size=13, bold=True, color="1B365D")
    ws4["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws4.row_dimensions[1].height = 26

    headers4 = [
        "Sl.No", "Coach Type", "Coach No", "Owning Rly", "Division", 
        "Year Built", "Age (Yrs)", "Age Classification", "Man-Hours", "Current Location / Pit"
    ]
    for col_idx, h in enumerate(headers4, 1):
        c = ws4.cell(row=3, column=col_idx, value=h)
        c.fill = blue_header
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border
    ws4.row_dimensions[3].height = 25

    all_sorted = sorted(coaches, key=lambda x: (x["group"], -(x["yb"] or 0), x["cno"]))
    row_t4 = 4
    sno_t4 = 1
    current_grp4 = None
    for c in all_sorted:
        if c["group"] != current_grp4:
            current_grp4 = c["group"]
            grp_ge = sum(1 for x in all_sorted if x["group"] == current_grp4 and x["is_ge2013"])
            grp_lt = sum(1 for x in all_sorted if x["group"] == current_grp4 and not x["is_ge2013"])
            grp_tot = grp_ge + grp_lt
            ws4.merge_cells(start_row=row_t4, start_column=1, end_row=row_t4, end_column=10)
            sh = ws4.cell(row=row_t4, column=1, value=f"{current_grp4} - TOTAL: {grp_tot} COACHES (>= 2013: {grp_ge} | < 2013: {grp_lt})")
            sh.fill = sub_fill
            sh.font = Font(name="Calibri", size=11, bold=True, color="1B365D")
            sh.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            for c_idx in range(1, 11):
                ws4.cell(row=row_t4, column=c_idx).border = thin_border
            ws4.row_dimensions[row_t4].height = 22
            row_t4 += 1

        is_ge = c["is_ge2013"]
        hrs_val = c["man_hours"] if c["man_hours"] is not None else "-"
        vals = [
            sno_t4, c["type"], c["cno"], c["rly"], c["div"], 
            c["yb"], c["age"], c["age_cat"], hrs_val, f"{c['pit']} ({c['area']})"
        ]
        for col_idx, v in enumerate(vals, 1):
            cell = ws4.cell(row=row_t4, column=col_idx, value=v)
            cell.font = data_font
            cell.border = thin_border
            if col_idx in (1, 2, 4, 5, 6, 7, 8, 9):
                cell.alignment = Alignment(horizontal="center")
            if col_idx == 3:
                cell.alignment = Alignment(horizontal="center")
                cell.font = data_bold
            if col_idx == 6:
                cell.fill = fill_ge2013 if is_ge else fill_lt2013
                cell.font = font_ge2013 if is_ge else font_lt2013
            elif col_idx == 8:
                cell.fill = fill_ge2013 if is_ge else fill_lt2013
                cell.font = font_ge2013 if is_ge else font_lt2013
            elif col_idx == 9 and isinstance(v, int):
                cell.number_format = "#,##0"
                cell.font = data_bold
        ws4.row_dimensions[row_t4].height = 20
        row_t4 += 1
        sno_t4 += 1

    ws4.cell(row=row_t4, column=2, value="TOTAL").font = total_font
    ws4.cell(row=row_t4, column=3, value=f"{tot_all} Coaches").font = total_font
    ws4.cell(row=row_t4, column=3).alignment = Alignment(horizontal="center")
    ws4.cell(row=row_t4, column=8, value=f">= 2013: {tot_ge} | < 2013: {tot_lt}").font = total_font
    ws4.cell(row=row_t4, column=8).alignment = Alignment(horizontal="center")
    
    tot_all_hrs = sum(c["man_hours"] for c in coaches if c["man_hours"] is not None)
    c_hrs4 = ws4.cell(row=row_t4, column=9, value=tot_all_hrs)
    c_hrs4.font = total_font
    c_hrs4.alignment = Alignment(horizontal="center")
    c_hrs4.number_format = "#,##0"

    for col_idx in range(1, 11):
        c = ws4.cell(row=row_t4, column=col_idx)
        c.fill = total_fill
        c.border = double_bottom_border
    ws4.row_dimensions[row_t4].height = 24

    ws4.column_dimensions["A"].width = 8
    ws4.column_dimensions["B"].width = 14
    ws4.column_dimensions["C"].width = 16
    ws4.column_dimensions["D"].width = 14
    ws4.column_dimensions["E"].width = 14
    ws4.column_dimensions["F"].width = 14
    ws4.column_dimensions["G"].width = 12
    ws4.column_dimensions["H"].width = 24
    ws4.column_dimensions["I"].width = 16
    ws4.column_dimensions["J"].width = 28

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()
