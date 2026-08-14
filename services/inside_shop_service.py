import os
import io
import json
import calendar
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from collections import Counter

# Centralized imports from ERP codebase
from services.erp_service import fetch_master, fetch_single, fetch_clean, _parse_date
from services.decoders import decode_family, decode_division, decode_repair

# Centralized Google Sheet config
CREDENTIALS_PATH = "D:\\JIJO\\information\\Coach Position\\credentials.json"
SHEET_KEY = "1xnJmZgqYODnJknNsas6J9k1f5NKp9h4x54nJp7AsOIY"


def get_rly_from_div(div_code):
    div_code = str(div_code).strip().upper()
    if "/" in div_code:
        parts = div_code.split("/")
        if len(parts) > 1 and parts[1]:
            return parts[1]
            
    rly_map = {
        "TVC": "SR", "MAS": "SR", "PGT": "SR", "SA": "SR", "TPJ": "SR", "MDU": "SR", "MAQ": "SR", "NCJ": "SR", "ED": "SR", "CBE": "SR", "SRR": "SR", "VM": "SR",
        "SC": "SCR", "HYB": "SCR", "BZA": "SCR", "GNT": "SCR", "GTL": "SCR", "NED": "SCR",
        "GHY": "NF", "APDJ": "NF", "LMG": "NF", "TSK": "NF", "RNY": "NF",
        "ADI": "WR", "BRC": "WR", "RTM": "WR", "BVP": "WR", "RJT": "WR", "MMCT": "WR",
        "HWH": "ER", "SDAH": "ER", "ASN": "ER", "MLDT": "ER",
        "MYS": "SWR", "SBC": "SWR", "UBL": "SWR",
        "TNP": "SR", "EDW": "SR", "MS": "SR", "JTJ": "SR"
    }
    return rly_map.get(div_code, "SR")


def map_coach_desc_to_code(desc, family):
    desc = str(desc).strip().upper()
    if family == "ICF":
        if "WGSCN" in desc or "SCN" in desc or "CN" in desc:
            return "CN"
        if "GSLRD" in desc or "GSLR" in desc:
            return "GSLRD"
        if "GSRD" in desc:
            return "GSRD"
        if "SLR" in desc:
            return "SLR"
        if "GS" in desc:
            return "GS"
        if "CZJ" in desc or "SCZJ" in desc:
            return "CZJ"
        if "CZRJ" in desc or "SCZRJ" in desc:
            return "CZRJ"
        if "CZ" in desc or "SCZ" in desc:
            return "CZ"
        if "VPU" in desc:
            return "VPU"
        if "VPH" in desc:
            return "VPH"
        if "ARMV" in desc:
            return "ARMV"
        if "ART" in desc:
            return "ART CONV"
    elif family == "LHB":
        if "LWACCW" in desc:
            return "LWACCW"
        if "LWACCN" in desc:
            return "LWACCN"
        if "LWCBAC" in desc:
            return "LWCBAC"
        if "LWSCN" in desc:
            return "LWSCN"
        if "LWS" in desc:
            return "LWS"
        if "LSLRD" in desc:
            return "LSLRD"
        if "LS5" in desc:
            return "LS5"
        if "LVPH" in desc:
            return "LVPH"
    elif family == "NMG":
        if "NMGHSR" in desc:
            return "NMGHSR CONV"
        if "NMGHS" in desc:
            return "NMGHS"
        if "NMG" in desc:
            return "NMG"
    elif family in ("DEMU", "MEMU", "EMU", "TW", "SPECIAL"):
        if "DEMUTC" in desc or "DEMU TC" in desc:
            return "DEMU TC"
        if "MEMUMC" in desc or "MEMU MC" in desc:
            return "MEMU MC"
        if "MEMUTC" in desc or "MEMU TC" in desc:
            return "MEMU TC"
        if "TW8W" in desc:
            return "TW8W"
        if "TW4W" in desc:
            return "TW4W"
        if "DPC" in desc:
            return "DPC"
        if "SPIC" in desc:
            return "SPIC"
        if "EMUTC" in desc or "EMU TC" in desc or "YSY" in desc or "YSD" in desc or "YFSY" in desc:
            return "EMU TC"
        if "EMUMC" in desc or "EMU MC" in desc or "DMSC" in desc or "YZZS" in desc:
            return "EMU MC"
        if "SPART" in desc:
            return "SPART"
        if desc == "TC":
            if family == "DEMU": return "DEMU TC"
            if family == "MEMU": return "MEMU TC"
            if family == "EMU": return "EMU TC"
    return None

def get_inside_shop_data():
    records = fetch_clean()
    
    icf_list = []
    lhb_list = []
    emu_list = []
    
    for rec in records:
        # Exclude coaches that have been physically despatched
        pitnum = str(rec.get("pitnum", "")).strip()
        if pitnum.upper() in ("DESP",):
            continue

        status = str(rec.get("status", "") or rec.get("pohstatus", "")).strip().upper()
        if status in ("DESPATCHED", "OUTTURN", "COMPLETED", "INACTIVE"):
            continue
            
        coachno = rec.get("coachno")
        coach_desc = str(rec.get("coach_desc") or rec.get("coachdesc") or "").strip().upper()
        demandid = rec.get("demandid")
        
        detail = {}
        if demandid:
            try:
                detail = fetch_single(demandid)
            except:
                pass
                
        # Exclude if it has a valid despatch date (outturned)
        desp_date = detail.get("desp_date") or detail.get("despdate") or ""
        if desp_date and str(desp_date).strip().lower() not in ("none", "null", "nan", ""):
            continue
            
        actual_desp = str(detail.get("actualdespdate") or "").strip()
        if actual_desp and actual_desp.lower() not in ("none", "null", "nan", ""):
            continue
            
        family = decode_family(coach_desc)
        division = decode_division(detail.get("dvnid") or rec.get("dvnid"))
        rly = get_rly_from_div(division)
        pdc = detail.get("pdc_date") or rec.get("pdc_date") or ""
        
        coach_item = {
            "rly": rly,
            "type": coach_desc,
            "coachno": coachno,
            "divn": division,
            "pdc": pdc
        }
        
        if family == "ICF":
            icf_list.append(coach_item)
        elif family == "LHB":
            lhb_list.append(coach_item)
        else:
            emu_list.append(coach_item)
            
    icf_counts = Counter(c["type"] for c in icf_list)
    lhb_counts = Counter(c["type"] for c in lhb_list)
    emu_counts = Counter(c["type"] for c in emu_list)
    
    return {
        "icf": icf_list,
        "icf_counts": dict(icf_counts),
        "lhb": lhb_list,
        "lhb_counts": dict(lhb_counts),
        "emu": emu_list,
        "emu_counts": dict(emu_counts)
    }

def generate_inside_shop_excel():
    data = get_inside_shop_data()
    icf = data["icf"]
    lhb = data["lhb"]
    emu = data["emu"]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shop Position"
    ws.views.sheetView[0].showGridLines = True
    
    font_title = Font(name="Calibri", size=14, bold=True, color="1B4F72")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    font_section = Font(name="Calibri", size=11, bold=True, color="1B4F72")
    font_data = Font(name="Calibri", size=10, bold=False)
    font_bold = Font(name="Calibri", size=10, bold=True)
    
    fill_header = PatternFill(start_color="1F618D", end_color="1F618D", fill_type="solid")
    fill_summary = PatternFill(start_color="D6EAF8", end_color="D6EAF8", fill_type="solid")
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='BDC3C7'), right=Side(style='thin', color='BDC3C7'),
        top=Side(style='thin', color='BDC3C7'), bottom=Side(style='thin', color='BDC3C7')
    )
    
    curr_row = 2
    
    # SECTION 1: ICF COACHES
    cell_title = ws.cell(row=curr_row, column=1, value="loco  in shop ICF")
    cell_title.font = font_title
    curr_row += 1
    
    headers = ["Sl No", "Rly", "Type", "Coach No.", "Divn", "PDC"]
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[curr_row].height = 20
    curr_row += 1
    
    for s_idx, c in enumerate(icf, 1):
        row_vals = [s_idx, c["rly"], c["type"], c["coachno"], c["divn"], c["pdc"]]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.font = font_bold if col_idx == 4 else font_data
            cell.alignment = align_center if col_idx in (1, 2, 3, 5, 6) else align_left
            cell.border = thin_border
        ws.row_dimensions[curr_row].height = 18
        curr_row += 1
        
    curr_row += 1
    cell_c_head1 = ws.cell(row=curr_row, column=1, value="TYPE")
    cell_c_head1.font = font_header
    cell_c_head1.fill = fill_header
    cell_c_head1.border = thin_border
    cell_c_head1.alignment = align_center
    
    cell_c_head2 = ws.cell(row=curr_row, column=2, value="Count of Coach No.")
    cell_c_head2.font = font_header
    cell_c_head2.fill = fill_header
    cell_c_head2.border = thin_border
    cell_c_head2.alignment = align_center
    curr_row += 1
    
    icf_counts = data["icf_counts"]
    for t, cnt in icf_counts.items():
        cell_t = ws.cell(row=curr_row, column=1, value=t)
        cell_t.font = font_bold
        cell_t.border = thin_border
        cell_t.alignment = align_center
        
        cell_cnt = ws.cell(row=curr_row, column=2, value=cnt)
        cell_cnt.font = font_data
        cell_cnt.border = thin_border
        cell_cnt.alignment = align_center
        curr_row += 1
        
    cell_gt = ws.cell(row=curr_row, column=1, value="Grand Total")
    cell_gt.font = font_bold
    cell_gt.fill = fill_summary
    cell_gt.border = thin_border
    cell_gt.alignment = align_center
    
    cell_gt_val = ws.cell(row=curr_row, column=2, value=len(icf))
    cell_gt_val.font = font_bold
    cell_gt_val.fill = fill_summary
    cell_gt_val.border = thin_border
    cell_gt_val.alignment = align_center
    curr_row += 3
    
    # SECTION 2: LHB COACHES
    cell_lhb_sec = ws.cell(row=curr_row, column=1, value="LHB-NAC/AC")
    cell_lhb_sec.font = font_section
    curr_row += 1
    
    cell_lhb_title = ws.cell(row=curr_row, column=1, value="loco  in shop LHB")
    cell_lhb_title.font = font_title
    curr_row += 1
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[curr_row].height = 20
    curr_row += 1
    
    for s_idx, c in enumerate(lhb, 1):
        row_vals = [s_idx, c["rly"], c["type"], c["coachno"], c["divn"], c["pdc"]]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.font = font_bold if col_idx == 4 else font_data
            cell.alignment = align_center if col_idx in (1, 2, 3, 5, 6) else align_left
            cell.border = thin_border
        ws.row_dimensions[curr_row].height = 18
        curr_row += 1
        
    curr_row += 1
    cell_lhb_head1 = ws.cell(row=curr_row, column=1, value="TYPE")
    cell_lhb_head1.font = font_header
    cell_lhb_head1.fill = fill_header
    cell_lhb_head1.border = thin_border
    cell_lhb_head1.alignment = align_center
    
    cell_lhb_head2 = ws.cell(row=curr_row, column=2, value="Count of Coach No.")
    cell_lhb_head2.font = font_header
    cell_lhb_head2.fill = fill_header
    cell_lhb_head2.border = thin_border
    cell_lhb_head2.alignment = align_center
    curr_row += 1
    
    lhb_counts = data["lhb_counts"]
    for t, cnt in lhb_counts.items():
        cell_t = ws.cell(row=curr_row, column=1, value=t)
        cell_t.font = font_bold
        cell_t.border = thin_border
        cell_t.alignment = align_center
        
        cell_cnt = ws.cell(row=curr_row, column=2, value=cnt)
        cell_cnt.font = font_data
        cell_cnt.border = thin_border
        cell_cnt.alignment = align_center
        curr_row += 1
        
    cell_lhb_gt = ws.cell(row=curr_row, column=1, value="Grand Total")
    cell_lhb_gt.font = font_bold
    cell_lhb_gt.fill = fill_summary
    cell_lhb_gt.border = thin_border
    cell_lhb_gt.alignment = align_center
    
    cell_lhb_gt_val = ws.cell(row=curr_row, column=2, value=len(lhb))
    cell_lhb_gt_val.font = font_bold
    cell_lhb_gt_val.fill = fill_summary
    cell_lhb_gt_val.border = thin_border
    cell_lhb_gt_val.alignment = align_center
    curr_row += 3
    
    # SECTION 3: EMU / MEMU / DEMU
    cell_emu_sec = ws.cell(row=curr_row, column=1, value="EMU-MC/TC")
    cell_emu_sec.font = font_section
    curr_row += 1
    
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=curr_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = thin_border
    ws.row_dimensions[curr_row].height = 20
    curr_row += 1
    
    for s_idx, c in enumerate(emu, 1):
        row_vals = [s_idx, c["rly"], c["type"], c["coachno"], c["divn"], c["pdc"]]
        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=curr_row, column=col_idx, value=val)
            cell.font = font_bold if col_idx == 4 else font_data
            cell.alignment = align_center if col_idx in (1, 2, 3, 5, 6) else align_left
            cell.border = thin_border
        ws.row_dimensions[curr_row].height = 18
        curr_row += 1

    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()


# ==============================================================================
# REPORT 5: TYPE-WISE HOLDING & POH PERFORMANCE REPORT (DAILY)
