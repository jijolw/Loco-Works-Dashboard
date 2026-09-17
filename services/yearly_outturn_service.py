import os
import io
import json
import calendar
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from services.type_wise_holding_service import get_type_wise_holding_report_data, STANDARD_WORKSHOP_CATEGORIES, fetch_live_keycloak_demands

def get_yearly_outturn_data(fy_str="2026-27", bypass_cache=False):
    """
    12-Month Outturn Matrix across any Financial Year derived directly from TypeWise Pure ERP pipeline.
    Runs months in parallel using ThreadPoolExecutor for instant response.
    """
    try:
        start_year = int(fy_str.split("-")[0])
    except:
        start_year = 2026
    end_year = start_year + 1
    
    fy_months = [
        ("April", start_year), ("May", start_year), ("June", start_year),
        ("July", start_year), ("August", start_year), ("September", start_year),
        ("October", start_year), ("November", start_year), ("December", start_year),
        ("January", end_year), ("February", end_year), ("March", end_year)
    ]
    
    months_labels = ["APR", "MAY", "JUN", "JUL", "AUG", "SEP", "OCT", "NOV", "DEC", "JAN", "FEB", "MAR"]
    grid = {}

    # Pre-populate demand cache once
    fetch_live_keycloak_demands()

    def fetch_month_outturn(item):
        m_name, y_val = item
        m_data = {cat: 0 for cat in STANDARD_WORKSHOP_CATEGORIES}
        try:
            rep = get_type_wise_holding_report_data(m_name, y_val)
            for r in rep["data"]:
                cat = r["coach_type"]
                tot = r["physical_despatch"] + r["fnd"]
                m_data[cat] = tot
        except Exception:
            pass
        return m_name, m_data

    with ThreadPoolExecutor(max_workers=6) as executor:
        results = executor.map(fetch_month_outturn, fy_months)
        for m_name, m_data in results:
            grid[m_name] = m_data
        
    return {
        "fy": fy_str,
        "months": [m[0] for m in fy_months],
        "months_labels": months_labels,
        "grid": grid,
        "categories": STANDARD_WORKSHOP_CATEGORIES
    }

def generate_yearly_outturn_excel(fy_str="2026-27"):
    data = get_yearly_outturn_data(fy_str)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Yearly_Outturn_{fy_str}"
    
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells("A1:N1")
    ws["A1"] = f"YEARLY OUTTURN SUMMARY (FY {fy_str}) - PURE ERP DATA"
    ws["A1"].font = Font(name="Arial", size=14, bold=True, color="1F497D")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    
    headers = ["Category"] + data["months_labels"] + ["Total"]
    ws.row_dimensions[3].height = 28
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    
    for c_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=c_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    row_idx = 4
    for cat in data["categories"]:
        row_tot = 0
        ws.cell(row=row_idx, column=1, value=cat).alignment = Alignment(horizontal="left", vertical="center")
        for c_idx, m_name in enumerate(data["months"], 2):
            val = data["grid"].get(m_name, {}).get(cat, 0)
            row_tot += val
            ws.cell(row=row_idx, column=c_idx, value=val).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row=row_idx, column=14, value=row_tot).alignment = Alignment(horizontal="center", vertical="center")
        row_idx += 1
        
    ws.cell(row=row_idx, column=1, value="TOTAL").font = Font(bold=True)
    for c_idx, m_name in enumerate(data["months"], 2):
        col_sum = sum(data["grid"].get(m_name, {}).get(cat, 0) for cat in data["categories"])
        ws.cell(row=row_idx, column=c_idx, value=col_sum).font = Font(bold=True)
        ws.cell(row=row_idx, column=c_idx).alignment = Alignment(horizontal="center", vertical="center")
    grand_tot = sum(sum(data["grid"].get(m, {}).get(cat, 0) for cat in data["categories"]) for m in data["months"])
    ws.cell(row=row_idx, column=14, value=grand_tot).font = Font(bold=True)
    ws.cell(row=row_idx, column=14).alignment = Alignment(horizontal="center", vertical="center")
    
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()

