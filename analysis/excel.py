"""
analysis/excel.py - plain Excel downloads (JIJO's format: merged centred title, grid,
auto-fit ignoring row 1, bold Coach Type + FND qty where FND > 0).
"""
import io

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

from core import rules as R

THIN = Side(style="thin", color="808080")


def _sheet(title, header):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(header))
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=12)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(header)
    return wb, ws


def _finish(wb, ws, name, bold_rows=()):
    n = ws.max_column
    for r in range(2, ws.max_row + 1):
        for c in range(1, n + 1):
            cell = ws.cell(r, c)
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
            if c > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            if r == 2 or r in bold_rows:
                cell.font = Font(bold=True)
    widths = {}
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(8, w + 2), 40)
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio, name


def month_position(pos, as_on):
    hdr = ["Coach Type", "Outturn", "Despatched", "FND", "Holding Yard", "Holding Work Area", "Total Holding"]
    wb, ws = _sheet(f"LW/PER - COACH POSITION {pos['month']:%B %Y} (as on {as_on:%d.%m.%Y})", hdr)
    for r in pos["rows"]:
        ws.append([r["group"], r["outturn"], r["desp"], r["fnd"], r["yard"], r["work"], r["holding"]])
        if r["fnd"]:
            ws.cell(ws.max_row, 1).font = Font(bold=True, size=11)
            ws.cell(ws.max_row, 4).font = Font(bold=True, size=11)
    t = pos["total"]
    ws.append(["TOTAL", t["outturn"], t["desp"], t["fnd"], t["yard"], t["work"], t["holding"]])
    bold = [ws.max_row]
    if pos["fnd_list"] or pos["prev_fnd"]:
        ws.append([])
        ws.append(["FND (outturned, not physically despatched)"])
        bold.append(ws.max_row)
        for d in pos["fnd_list"] + pos["prev_fnd"]:
            ws.append([d["_group"], d["coachNo"], R.fmt(R.pdate(d["dispatchDate"]))])
    return _finish(wb, ws, f"Month_Position_{pos['month']:%b_%Y}.xlsx", bold)


def holding(h, as_on):
    hdr = ["Coach Type", "Yard", "Work Area", "Total", "0-20 days", "21-30 days", ">30 days",
           "Light (<500 h)", "Heavy (>=500 h)", "Max days"]
    wb, ws = _sheet(f"LW/PER - TYPE-WISE POH HOLDING as on {as_on:%d.%m.%Y}", hdr)
    for r in h["rows"]:
        ws.append([r["group"], r["yard"], r["work"], r["total"], r["0-20"], r["21-30"], r[">30"],
                   r["LIGHT"], r["HEAVY"], r["max"]])
    t = h["total"]
    ws.append(["TOTAL", t["yard"], t["work"], t["total"], t["0-20"], t["21-30"], t[">30"],
               t["LIGHT"], t["HEAVY"], t["max"]])
    return _finish(wb, ws, f"TypeWise_Holding_{as_on:%d%m%Y}.xlsx", [ws.max_row])


def outturn(x, as_on):
    hdr = ["Coach Type"] + [f"{m:%b-%y}" for m in x["months"]] + ["Total"]
    wb, ws = _sheet(f"LW/PER - MONTH-WISE OUTTURN {x['fy']} (desp date; as on {as_on:%d.%m.%Y})", hdr)
    for r in x["rows"]:
        ws.append([r["group"]] + r["vals"] + [r["total"]])
    ws.append(["TOTAL"] + x["totals"] + [x["grand"]])
    return _finish(wb, ws, f"Outturn_Monthwise_{x['fy']}.xlsx", [ws.max_row])


def _fmt(d):
    return d.strftime("%d.%m.%Y") if d else ""


def cycle(c, d0, d1, as_on):
    hdr = ["Coach Type", "Outturns", "Avg days entry->outturn", "Median", "Min", "Max", ">30 days",
           "Avg TFR->outturn", "Avg entry->physical desp", "Light avg", "Heavy avg"]
    wb, ws = _sheet(f"LW/PER - POH CYCLE TIME ({_fmt(d0)} to {_fmt(d1)}) - outturn by desp date, "
                    f"return/condemned excluded", hdr)
    for r in c["rows"] + ([dict(group="TOTAL", **c["total"])] if c["total"] else []):
        ws.append([r["group"], r["n"], r["avg"], r["med"], r["min"], r["max"], r["over30"], r["attn"],
                   r["physical"], r["light"], r["heavy"]])
    out = _finish(wb, ws, "", [ws.max_row])[0]
    wb2 = openpyxl.load_workbook(out)
    ws2 = wb2.create_sheet("Coach-wise")
    ws2.append(["Coach Type", "Coach No", "Received", "TFR", "Outturn (desp)", "Physical desp", "Days",
                "TFR->outturn", "Corr start", "Corr comp", "Corr days", "Pre hrs", "Final hrs"])
    for i in c["items"]:
        ws2.append([i["group"], i["coach"], _fmt(i["recd"]), _fmt(i["tfr"]), _fmt(i["desp"]), _fmt(i["actual"]),
                    i["total"], i["attn"], _fmt(i["cs"]), _fmt(i["cc"]), i["corr"], i["pre"], i["final"]])
    for col in range(1, 14):
        ws2.column_dimensions[get_column_letter(col)].width = 13
    bio = io.BytesIO(); wb2.save(bio); bio.seek(0)
    return bio, f"POH_Cycle_Time_{d0:%d%m%Y}_to_{d1:%d%m%Y}.xlsx"


def corrosion(k, d0, d1, as_on):
    hdr = ["Coach Type", "Outturns", "With corr dates", "Avg corrosion days", "Median", "Max",
           "Entry->corr start", "Corr done->outturn"]
    wb, ws = _sheet(f"LW/PER - CORROSION CYCLE TIME ({_fmt(d0)} to {_fmt(d1)})", hdr)
    for r in k["rows"] + ([dict(group="TOTAL", **k["total"])] if k["total"] else []):
        ws.append([r["group"], r["n"], r["with_corr"], r["corr"], r["corr_med"], r["corr_max"], r["to_corr"], r["after"]])
    bold = [ws.max_row]
    ws.append([])
    ws.append([f"{k['month']:%B %Y}", "Completed this month", "Carry forward", "Under corrosion now"])
    bold.append(ws.max_row)
    for r in k["month_rows"]:
        ws.append([r["group"], r["done"], r["carry"], r["under"]])
    return _finish(wb, ws, f"Corrosion_Cycle_{d0:%d%m%Y}_to_{d1:%d%m%Y}.xlsx", bold)


def hours(h, d0, d1, as_on):
    hdr = ["Coach Type", "Outturns", "Pre entered", "Avg pre hrs", "Final entered", "Avg final hrs",
           "Final vs pre %", "Light <500", "Heavy >=500", "Final hrs missing"]
    wb, ws = _sheet(f"LW/PER - MAN-HOURS ({_fmt(d0)} to {_fmt(d1)})", hdr)
    for r in h["rows"] + ([dict(group="TOTAL", **h["total"])] if h["total"] else []):
        ws.append([r["group"], r["n"], r["pre_n"], r["pre"], r["final_n"], r["final"], r["var"], r["light"],
                   r["heavy"], r["no_final"]])
    bold = [ws.max_row]
    ws.append([])
    ws.append(["Final hours missing", "Coach No", "Outturn", "Pre hrs"])
    bold.append(ws.max_row)
    for i in h["missing"]:
        ws.append([i["group"], i["coach"], _fmt(i["desp"]), i["pre"]])
    return _finish(wb, ws, f"Man_Hours_{d0:%d%m%Y}_to_{d1:%d%m%Y}.xlsx", bold)


def prev_poh(p, d0, d1, as_on):
    hdr = ["Last POH workshop", "Coaches", "Avg pre hrs", "Heavy %", "Avg final hrs", "Avg cycle days",
           "Avg months since last POH"]
    wb, ws = _sheet(f"LW/PER - PREVIOUS POH ANALYSIS ({_fmt(d0)} to {_fmt(d1)})", hdr)
    for r in p["rows"]:
        ws.append([r["wks"], r["n"], r["pre"], r["heavy_pct"], r["final"], r["cycle"], r["gap"]])
    bold = []
    if p["sched"]:
        ws.append([])
        ws.append(["LHB schedule", "SS1", "SS2", "SS3", "Other"])
        bold.append(ws.max_row)
        for g, s in p["sched"].items():
            ws.append([g] + list(s.values()))
    return _finish(wb, ws, f"Previous_POH_{d0:%d%m%Y}_to_{d1:%d%m%Y}.xlsx", bold)


def targets(t, as_on, s=None):
    hdr = ["Target type"]
    for m in t["months"]:
        hdr += [f"{m:%b-%y} T", f"{m:%b-%y} A"]
    hdr += ["Target (months set)", "Actual (same months)", "%", "Actual all months"]
    wb, ws = _sheet(f"LW/PER - TARGET vs ACTUAL {t['fy']} (as on {as_on:%d.%m.%Y}; '-' = target not set)", hdr)
    for r in t["rows"]:
        row = [r["type"]]
        for c in r["cells"]:
            row += ["-" if c["t"] is None else c["t"], c["a"]]
        ws.append(row + [r["tsum"], r["asum_t"], r["pct"], r["asum"]])
        for p in r.get("parts", []):
            prow = ["   " + p["name"]]
            for x in p["vals"]:
                prow += ["", x]
            ws.append(prow + ["", "", "", p["total"]])
    out = _finish(wb, ws, "", [])[0]
    if s:
        wb = openpyxl.load_workbook(out)
        ws = wb.create_sheet("By stock type", 0)
        ws.append([f"LW/PER - TARGET vs ACTUAL BY STOCK TYPE {s['fy']} (as on {as_on:%d.%m.%Y}; n/e = not in coach ERP)"])
        ws.append(["Stock type", "HQ T to date", "Actual (HQ months)", "HQ %", "Internal T", "Actual", "Internal %",
                   "RB annual", "RB pro-rata", "RB %", "Actual FY"])
        for r in s["rows"]:
            ne = "n/e"
            ws.append([r["label"], r["hq_t"], r["hq_a"] if r["in_erp"] else ne, r["hq_pct"], r["i_t"], r["i_a"],
                       r["i_pct"], r["rb_annual"], r["rb_todate"], r["rb_pct"], r["a_all"] if r["in_erp"] else ne])
        for col in range(1, 12):
            ws.column_dimensions[get_column_letter(col)].width = 16 if col > 1 else 26
        out = io.BytesIO(); wb.save(out); out.seek(0)
    return out, f"Target_vs_Actual_{t['fy']}.xlsx"
